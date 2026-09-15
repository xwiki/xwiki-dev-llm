"""Slide primitives: text, bullets, boxes, pictures, titles, per-slide chrome, speaker notes.

Every helper takes inches from the top-left and returns the shape it made, so a build script
reads as a layout rather than as pptx plumbing.
"""
import os

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from . import theme as t


def textbox(slide, x, y, w, h, text, size=18, color=None, bold=False, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, italic=False, space_after=6, line=None):
    """A text box with the deck's font applied to every run.

    `text` is a string (newlines become paragraphs) or a sequence of lines. Margins are zeroed
    so the box's x/y is where the glyphs actually start — without that, aligning a caption to
    the column above it is guesswork.
    """
    color = color or t.INK
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text.split('\n') if isinstance(text, str) else list(text)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line:
            p.line_spacing = line
        r = p.add_run()
        r.text = ln
        f = r.font
        f.name, f.size, f.bold, f.italic = t.FONT, Pt(size), bold, italic
        f.color.rgb = t.rgb(color)
    return tb


def bullets(slide, x, y, w, h, items, size=17, gap=9, line=1.05):
    """A bullet list.

    Each item is a string, or a `(text, bold, colour)` tuple, or `(text, bold, colour, indent)`
    where `indent` is in inches. The indent is a real paragraph indent rather than leading
    spaces, so a wrapped continuation line stays aligned under the first — leading spaces look
    identical until a line wraps, then break.
    """
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        indent = 0
        if isinstance(it, tuple):
            txt, bold, col = it[:3]
            if len(it) > 3:
                indent = it[3]
        else:
            txt, bold, col = it, False, t.INK
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = line
        if indent:
            pPr = p._p.get_or_add_pPr()
            pPr.set('marL', str(Inches(indent)))
            pPr.set('indent', '0')
        r = p.add_run()
        r.text = txt
        f = r.font
        f.name, f.size, f.bold = t.FONT, Pt(size), bold
        f.color.rgb = t.rgb(col)
    return tb


def text_height(items, w, size, gap=9, line=1.05, pad=0.0):
    """Estimate, in inches, the height `bullets()` will need for `items` in a box `w` wide.

    A box is drawn at whatever height it is given, so text that does not fit simply spills past
    it and out of any card behind it — silently, and only visible once the deck is rendered.
    Sizing the card from its content is what prevents that.

    The estimate is deliberately slightly generous. `0.47 * size` is the average glyph advance
    of a humanist sans as a fraction of its point size; a line of full-width capitals will
    exceed it, so leave a little room for a heading that might.
    """
    chars_per_line = max(8, int(w * 72 / (0.47 * size)))
    lines = 0
    for it in items:
        txt = it[0] if isinstance(it, tuple) else it
        lines += max(1, -(-len(txt) // chars_per_line))
    return lines * size * line * 1.2 / 72 + len(items) * gap / 72 + pad


def rect(slide, x, y, w, h, fill=None, outline=None, line_w=1.0, radius=0.04):
    """A rounded rectangle, used as a card behind grouped content.

    Shadows are switched off explicitly: pptx inherits a theme shadow otherwise, which survives
    into the PDF as a grey smear nobody asked for.
    """
    outline = t.RULE if outline is None else outline
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    sh.adjustments[0] = radius
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = t.rgb(fill)
    if outline is False:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = t.rgb(outline)
        sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    sh.text_frame.text = ''
    return sh


def picture(slide, path, x, y, max_w, max_h, center=True):
    """Place a picture scaled to fit the box without distorting it.

    python-pptx stretches a picture to whatever width and height it is given, so the aspect
    ratio has to be preserved here. 96 dpi is the reference the pptx geometry assumes.
    """
    from PIL import Image
    iw, ih = Image.open(path).size
    scale = min(max_w / (iw / 96.0), max_h / (ih / 96.0))
    w, h = (iw / 96.0) * scale, (ih / 96.0) * scale
    px = x + (max_w - w) / 2 if center else x
    py = y + (max_h - h) / 2 if center else y
    return slide.shapes.add_picture(path, Inches(px), Inches(py), Inches(w), Inches(h))


def blank(prs):
    """A slide with no placeholders. Layout 6 is the blank layout in the default template."""
    return prs.slides.add_slide(prs.slide_layouts[6])


def title(slide, text, sub=None, size=34):
    """Centred bold title with an optional muted one-line subtitle under it."""
    textbox(slide, t.MARGIN, 0.30, t.W - 2 * t.MARGIN, 0.72, text, size=size, bold=True,
            align=PP_ALIGN.CENTER)
    if sub:
        textbox(slide, t.MARGIN, 1.02, t.W - 2 * t.MARGIN, 0.34, sub, size=13, color=t.MUTED,
                align=PP_ALIGN.CENTER)


def chrome(slide, byline='', logo=None, badge=None, marker=None):
    """The footer every slide carries: byline left, optional marker, badge and logo right.

    `badge` is a licence image (e.g. a Creative Commons badge) for a deck meant to be
    published; `marker` is a short word such as 'Internal' for one that is not. Pass one or
    neither — a deck that claims both a public licence and confidentiality says nothing.
    """
    if byline:
        textbox(slide, t.MARGIN, t.H - 0.42, 4.2, 0.28, byline, size=9, color=t.MUTED)
    if marker:
        textbox(slide, t.W / 2 - 1.0, t.H - 0.42, 2.0, 0.28, marker, size=9, color=t.MUTED,
                bold=True, align=PP_ALIGN.CENTER)
    if logo and os.path.exists(logo):
        picture(slide, logo, t.W - t.MARGIN - 0.95, t.H - 0.62, 0.95, 0.46, center=False)
    if badge and os.path.exists(badge):
        picture(slide, badge, t.W - t.MARGIN - 2.15, t.H - 0.58, 1.05, 0.38, center=False)


def notes(slide, text, minutes):
    """Attach speaker notes, prefixed with this slide's minute budget.

    The `[n min]` prefix is the contract checks.py and notes.py both read: it is what lets the
    deck report its own running time instead of the author guessing it.
    """
    slide.notes_slide.notes_text_frame.text = f'[{minutes} min]\n\n' + text.strip()


def delta_chip(slide, x, y, w, value, good_up=True, size=11):
    """A small up/down chip whose colour is semantic: green means the direction wanted.

    `value` is a percentage, or a `(sign, label)` pair when the change reads better as an
    absolute or a multiplier — `(+1, '+1')`, `(+1, 'x11')` — because a growth of 1000% is
    legible as a multiplier and not as a percentage.
    """
    if value is None:
        return None
    if isinstance(value, tuple):
        sign, label = value
        up = sign >= 0
    else:
        up = value >= 0
        label = f'{abs(value):.0f}%'
    col = t.GREEN if up == good_up else t.RED
    arrow = '↑' if up else '↓'
    return textbox(slide, x, y, w, 0.24, f'{arrow} {label}', size=size, color=col, bold=True)
